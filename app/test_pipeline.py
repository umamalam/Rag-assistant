"""
test_pipeline.py -- offline tests (no model download, no API calls).
Uses a hashed bag-of-words stand-in for the embedder and a fake LLM.
Run:  python3 test_pipeline.py
"""

import hashlib
import math
import os
import re
import sys
import tempfile

tmp = tempfile.mkdtemp()
os.environ["DB_DIR"] = os.path.join(tmp, "chroma")
os.environ["DOCUMENTS_DIR"] = os.path.join(tmp, "docs")
os.environ["LLM_PROVIDER"] = "groq"

import ingest, llm, loaders, privacy, rag_engine, store  # noqa: E402


def fake_embed(texts):
    out = []
    for t in texts:
        v = [0.0] * 512
        for w in re.findall(r"[a-z]+", t.lower()):
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % 512] += 1
        n = math.sqrt(sum(x * x for x in v)) or 1
        out.append([x / n for x in v])
    return out


store.embed = fake_embed
sent = []  # everything the "LLM" receives
def fake_chat(messages, temperature=0.2):
    sent.append(messages)
    if messages[0]["content"].startswith("Rewrite"):
        return "sharding rollout plan"  # a plausible standalone rewrite
    return "ANSWER [1]"


llm.chat = fake_chat

docs = os.environ["DOCUMENTS_DIR"]
os.makedirs(docs)


def write(name, text):
    with open(os.path.join(docs, name), "w") as f:
        f.write(text)


checks = []


def check(name, cond):
    checks.append(cond)
    print(("PASS  " if cond else "FAIL  ") + name)


# ---- fixtures: markdown, txt, eml, mbox, docx, pdf ----
write("sysdesign.md", "# Caching\nCaching stores hot data in memory to cut database load. "
      "Use a cache aside pattern with eviction policies.\n# Sharding\nSharding splits a "
      "database across many machines by key to scale writes horizontally.\n")
write("notes.txt", "Meeting notes: we decided to adopt sharding for the orders database "
      "next quarter because write load keeps growing.")
write("mail.eml", "From: alice@example.com\nTo: me@example.com\nSubject: Sharding plan\n"
      "Date: Mon, 01 Jan 2024 10:00:00 +0000\nContent-Type: text/plain\n\n"
      "Hi, call me on +1 415 555 0132 about the sharding rollout. Card 4111 1111 1111 1111.\n")
write("inbox.mbox",
      "From a@x.com Mon Jan  1 10:00:00 2024\nFrom: a@x.com\nSubject: One\nDate: Mon, 01 Jan 2024 10:00:00 +0000\n\nFirst message about lunch plans.\n\n"
      "From b@x.com Tue Jan  2 10:00:00 2024\nFrom: b@x.com\nSubject: Two\nDate: Tue, 02 Jan 2024 10:00:00 +0000\n\nSecond message about travel plans.\n\n")

from docx import Document  # noqa: E402
d = Document(); d.add_heading("Kubernetes", 1); d.add_paragraph("Pods are scheduled onto nodes by the scheduler.")
d.save(os.path.join(docs, "k8s.docx"))

from reportlab.pdfgen import canvas  # noqa: E402
c = canvas.Canvas(os.path.join(docs, "book.pdf"))
for p in range(1, 41):
    c.drawString(72, 750, f"Chapter {p}: this page discusses topic{p} and load balancing basics number {p}.")
    c.showPage()
c.save()

# ---- loaders ----
check("markdown splits by heading", [u["location"] for u in loaders.load_file(f"{docs}/sysdesign.md")]
      == ["section: Caching", "section: Sharding"])
check("docx heading section", loaders.load_file(f"{docs}/k8s.docx")[0]["location"] == "section: Kubernetes")
check("eml parsed with date", "doc_ts" in loaders.load_file(f"{docs}/mail.eml")[0])
check("mbox yields 2 messages", len(loaders.load_file(f"{docs}/inbox.mbox")) == 2)
check("pdf yields 40 pages", len(loaders.load_file(f"{docs}/book.pdf")) == 40)

# ---- chunker ----
chunks = ingest.chunk_text("Sentence number one is here. " * 100)
check("chunks bounded", all(len(x) <= 800 for x in chunks) and len(chunks) > 1)
check("no redundant tail chunk", len(ingest.chunk_text("a" * 900 + " word.")) == 2)

# ---- ingest: first run, skip, change, remove ----
col = store.get_collection()
n1 = sum(ingest.ingest_file(col, docs, p) for p in ingest.find_files(docs))
check("all 6 file types indexed", len(rag_engine.list_sources()) == 6)
n2 = sum(ingest.ingest_file(col, docs, p) for p in ingest.find_files(docs))
check("re-run skips unchanged", n2 == 0)
write("notes.txt", "Meeting notes: we decided to adopt CACHING everywhere instead.")
check("changed file re-indexed", ingest.ingest_file(col, docs, os.path.join(docs, "notes.txt")) > 0)
check("stale text gone", not any("orders database" in d for d in col.get(where={"source": "notes.txt"})["documents"]))
col.delete(where={"source": "notes.txt"})
check("remove works", "notes.txt" not in rag_engine.list_sources())

# ---- retrieval quality ----
r = rag_engine.retrieve("How does sharding scale writes?")
check("relevant source ranks first", r and r[0]["source"] in ("sysdesign.md", "mail.eml"))
check("relevance is cosine (0..1)", all(0 <= c["relevance"] <= 1 for c in r))
cands = [{"id": f"a{i}", "source": "A"} for i in range(5)] + [{"id": "b0", "source": "B"}, {"id": "c0", "source": "C"}]
picked = rag_engine._diversify(cands, top_k=4, cap=2)
check("diversify: cap of 2 lets B and C in", [c["id"] for c in picked] == ["a0", "a1", "b0", "c0"])
check("diversify: backfills when few sources", len(rag_engine._diversify(cands[:5], top_k=4, cap=2)) == 4)
check("single-source filter lifts the cap", len(rag_engine.retrieve("load balancing", top_k=6, sources=["book.pdf"])) == 6)
check("date filter excludes old docs", rag_engine.retrieve("sharding", since_ts=2_000_000_000) == [])
check("irrelevant question -> no chunks", rag_engine.retrieve("zebra quantum xylophone") == [])

# ---- ask(): no-info short-circuit, memory, privacy ----
sent.clear()
a = rag_engine.ask("zebra quantum xylophone")
check("out-of-scope answers without calling LLM", a["answer"] == rag_engine.NO_INFO and not sent)

sent.clear()
rag_engine.ask("What did the email say about sharding?")
payload = str(sent[-1])
check("emails/phones/cards redacted before cloud", "alice@example.com" not in payload
      and "4111" not in payload and "555 0132" not in payload)

os.environ["REDACT_PII"] = "0"
sent.clear()
rag_engine.ask("What did the email say about sharding?")
check("redaction can be disabled", "alice@example.com" in str(sent[-1]))
os.environ.pop("REDACT_PII")

sent.clear()
rag_engine.ask("and the rollout?", history=[{"role": "user", "content": "sharding plan"},
                                            {"role": "assistant", "content": "It is planned."}])
check("follow-up triggers rewrite call", len(sent) == 2)

check("injection warning in system prompt", "untrusted" in rag_engine.SYSTEM_PROMPT)

# ---- summarize ----
sent.clear()
s = rag_engine.summarize("book.pdf")
check("summary samples <= 16 chunks across the doc", len(s["sources"]) == 16
      and s["sources"][0]["page"] == 1 and s["sources"][-1]["page"] == 40)
check("summarize unknown file handled", "not in the index" in rag_engine.summarize("nope.pdf")["answer"])

# ---- privacy module ----
check("redact secrets", "[SECRET]" in privacy.redact("key " + "ghp_" + "a" * 36))

# ---- large batch (the old code could fail here) ----
big = " ".join(f"Sentence {i} about batching." for i in range(6000))
write("big.txt", big)
check("very large file indexes in batches", ingest.ingest_file(col, docs, os.path.join(docs, "big.txt")) > ingest.BATCH)

print(f"\n{sum(checks)}/{len(checks)} checks passed")
sys.exit(0 if all(checks) else 1)
